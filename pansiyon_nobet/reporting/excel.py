from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


def _apply_sheet_formatting(ws: Worksheet) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = True

    # Yazdırma ayarları (yatay)
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4

    # Kolon genişlikleri
    widths = [12, 12, 55, 55]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    header_font = Font(bold=True)
    center = Alignment(vertical="center", wrap_text=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = center

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = center


def write_monthly_xlsx(*, rows: list[dict[str, str]], output_path: str) -> str:
    """
    Aylık nöbet listesini Excel olarak yazar.
    """
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Aylık Nöbet"

    headers = ["Tarih", "Gün", "Erkek Pansiyon", "Kız Pansiyon"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])

    _apply_sheet_formatting(ws)
    wb.save(out)
    return str(out)

